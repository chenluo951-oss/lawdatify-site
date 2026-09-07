#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""《合规资讯周报》Word 生成器（五章节，标宋大标题/黑体小标题/楷体正文）"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
_DATAMOD = sys.argv[1] if len(sys.argv) > 1 else "weekly_data"
WD = __import__(_DATAMOD)
import gen_report_pdf as G  # 复用 prev_issue_dates 推算往期回顾
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_DIR = os.environ.get("CBR_OUT_DIR") or "/Users/luochen/Desktop/合规资讯简报"

NAVY = (0x1F, 0x3B, 0x63)
NAVY2 = (0x2E, 0x5E, 0x8C)
GRAY = (0x66, 0x66, 0x66)
LGRAY = (0x99, 0x99, 0x99)
GOLD = (0x8A, 0x4B, 0x08)

SONG = "方正小标宋简体"
HEITI = "微软雅黑"
KAITI = "楷体"


def set_font(run, size=10.5, bold=False, color=None, name=KAITI):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    r = run._element.get_or_add_rPr()
    rf = r.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        r.append(rf)
    rf.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_h1(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_font(run, 15.5, True, NAVY, SONG)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_h2(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_font(run, 13, True, NAVY2, HEITI)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    return p


def add_body(doc, text, color=None, bold=False, name=KAITI, size=10.5):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_font(run, size, bold, color, name)
    p.paragraph_format.line_spacing = 1.5
    return p


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def main():
    M = WD.META
    D = WD.DATA
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # 封面
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(M["title"]), 22, True, NAVY, SONG)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(M["brief_en"]), 11, False, LGRAY, HEITI)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(M["date_str"]), 13, False, GRAY, KAITI)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(M["subtitle"]), 10, False, LGRAY, HEITI)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("朴朴超市 · 法务合规部"), 11, True, NAVY, HEITI)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("守护合规底线 · 支撑业务决策"), 9, False, GRAY, HEITI)
    n_policy = sum(len(items) for _, items in D["policy"])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(f"本期导读：六大领域重点动态 {n_policy} 项 · 通报处罚 {len(D['penalties'])} 起 · "
                       f"朴朴专题 {len(D['pupu_items'])} 项 · 下周前瞻 {len(D['outlook'])} 条"),
             9.5, False, GRAY, HEITI)
    prev = G.prev_issue_dates(M["date_str"], n=3)
    if prev:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run("近 3 期回顾：　" + "　·　".join(prev)), 9, False, LGRAY, HEITI)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("覆盖口径：六大合规领域公开监管动态与通报处罚"), 9, False, GRAY, HEITI)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("审核与签发：朴朴超市法务合规部"), 9, False, GRAY, HEITI)
    doc.add_paragraph()

    # 目录
    add_h1(doc, "目　录")
    for sec in M["sections"].values():
        add_body(doc, sec, color=NAVY2, bold=True, name=HEITI, size=12)

    # 一、本周综述
    add_h1(doc, M["sections"]["summary"])
    for i, s in enumerate(D["summary"], 1):
        add_body(doc, f"{i}. {s}")

    # 二、六大领域动态回顾
    add_h1(doc, M["sections"]["policy"])
    for di, (domain, items) in enumerate(D["policy"], 1):
        add_h2(doc, f"{di}. {domain}")
        for it in items:
            p = doc.add_paragraph()
            set_font(p.add_run(f"▍{it['title']}"), 11.5, True, None, HEITI)
            add_body(doc, it["meta"], color=LGRAY, name=HEITI, size=9)
            if it.get("url"):
                add_body(doc, f"原文链接：{it['url']}", color=NAVY2, name=KAITI, size=9)
            add_body(doc, f"要点：{it['content']}")
            add_body(doc, f"解读：{it['analysis']}", color=GOLD)
        doc.add_paragraph()

    # 三、监管通报与处罚汇总
    add_h1(doc, M["sections"]["penalties"])
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, t in enumerate(["时间", "监管机构", "涉及对象", "违规事由", "处置措施", "来源"]):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(t), 10, True, (255, 255, 255), HEITI)
        shade_cell(hdr[i], "1F3B63")
    for row_data in D["penalties"]:
        row = table.add_row().cells
        for i, v in enumerate(row_data):
            row[i].text = ""
            set_font(row[i].paragraphs[0].add_run(v), 9, False, None, KAITI)
    doc.add_paragraph()
    add_h2(doc, "本周处罚统计小结")
    for i, s in enumerate(D["penalty_stats"], 1):
        add_body(doc, f"{i}. {s}")

    # 四、朴朴超市业务专题
    add_h1(doc, M["sections"]["pupu"])
    add_body(doc, "说明：本板块聚焦与朴朴超市（生鲜电商/前置仓即时零售）业务直接相关的监管动态，"
                  "逐项评估合规风险等级并给出应对建议。", color=GRAY, name=HEITI, size=10)
    risk_color = {"高": (0xC0, 0x39, 0x2B), "中高": (0xD3, 0x54, 0x00),
                  "中": (0xB8, 0x86, 0x0B), "低": (0x5B, 0x8C, 0x5A)}
    for title, risk, scope, impact, action in D["pupu_items"]:
        p = doc.add_paragraph()
        set_font(p.add_run(f"▍{title}"), 11.5, True, None, HEITI)
        set_font(p.add_run(f"  【风险等级：{risk}】"), 10.5, True, risk_color.get(risk, GRAY), HEITI)
        add_body(doc, f"涉及业务环节：{scope}", color=NAVY2, bold=True, name=HEITI, size=10)
        add_body(doc, f"影响分析：{impact}")
        add_body(doc, f"应对建议：{action}", color=GOLD)
    doc.add_paragraph()

    # 五、下周前瞻
    add_h1(doc, M["sections"]["outlook"])
    for i, r in enumerate(D["outlook"], 1):
        add_body(doc, f"{i}. {r}")

    # 信息来源与免责声明（固定尾部栏）
    doc.add_paragraph()
    add_h2(doc, "信息来源与免责声明")
    for _t in [
        "一、信息来源：本周报内容基于公开网络检索整理的监管机构官方网站通报、公告、处罚决定书及权威媒体报道，各条动态均附原文链接，便于溯源核验。",
        "二、用途限制：本周报仅供朴朴超市内部合规参考，不构成法律意见或决策依据；据此采取具体合规措施前，请咨询法务或外部专业律师。",
        "三、时效性：监管政策与执法动态更新较快，本周报截至生成时点整理，不排除后续修订或新发文件导致内容变化，请以最新官方发布为准。",
        "四、版权：本周报由自动化任务生成，引用内容著作权归原发布机构所有。",
    ]:
        add_body(doc, _t, color=GRAY, size=9.5)

    # 尾注
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("— 本周报由自动化任务每周生成，内容基于公开网络信息检索整理，仅供参考 —"),
             9, False, LGRAY, HEITI)

    # 页脚：左来源/免责声明 + 右页码（对齐 PDF footer，benchmark C③）
    G.add_docx_footer(doc)

    os.makedirs(OUT_DIR, exist_ok=True)
    docx_path = os.path.join(OUT_DIR, M["filename"].replace(".pdf", ".docx"))
    doc.save(docx_path)
    print(f"DOCX saved: {docx_path}")
    return docx_path


if __name__ == "__main__":
    main()
