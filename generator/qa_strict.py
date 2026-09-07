#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""合规资讯报告 —— 交付前综合质量门禁（严格版）

检测项：
  P0-1 缺字检测：报告数据字符集 ⊆ PDF 提取字符集（差集必须为空）
  P0-2 tofu 检测：PDF 中不得出现 U+FFFD 替换字符
  P0-3 链接深链：所有官方链接必须为具体页面，禁止官网首页根域名（beian 除外）
  P1-1 章节齐全
  P1-2 图示内嵌（深度版）
  P1-3 页数达标（深度版 ≥30）
"""
import sys, os, re
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

DATAMOD = sys.argv[1] if len(sys.argv) > 1 else "daily_data_0901"
MOD = __import__(DATAMOD)
import deep_data as DD
from pypdf import PdfReader

OUT_DIR = os.environ.get("CBR_OUT_DIR") or "/Users/luochen/Desktop/合规资讯简报"

# ---------- 1. 收集报告实际用到的全部字符 ----------
def collect_chars(include_deep=False):
    M, D = MOD.META, MOD.DATA
    parts = [M["title"], M["date_str"], M["header_text"], M["subtitle"],
             M["brief_en"], M["filename"], M["tagline"], "目　录",
             "本期导读", "处罚统计小结", "每周监管与合规动态", "每日监管与合规动态"]
    parts.extend(M["sections"].values())
    parts.extend(D["summary"])
    for domain, items in D["policy"]:
        parts.append(domain)
        for it in items:
            parts.extend([it["title"], it["meta"], it["content"], it["analysis"], it.get("url", "")])
    for r in D["penalties"]:
        parts.extend(r)
    parts.extend(D["penalty_stats"])
    for t in D["pupu_items"]:
        parts.extend(t)
    parts.extend(D["outlook"])
    for row in D["matrix_rows"]:
        parts.extend(row)
    # 界面固定文案（两版共用）
    parts.extend(["风险主题", "采购", "仓储/加工", "线上运营", "配送", "会员营销", "用工",
                  "高", "中高", "中", "低", "—", "【要点】", "涉及业务环节：", "影响分析：",
                  "应对建议：", "解读：", "原文链接：", "来源", "时间", "监管机构",
                  "涉及对象", "违规事由", "处置措施", "第", "页", "共"])
    if not include_deep:
        return "".join(parts)
    # ---- 深度版叠加的共享内容（含各表表头，历史漏收集曾致缺字）----
    for k, v in DD.DOMAIN_DEEP.items():
        parts.append(k); parts.extend(v.keys()); parts.extend(v.values())
    parts.extend(DD.INDUSTRY_SURVEY["cols"])
    for row in DD.INDUSTRY_SURVEY["rows"]:
        parts.extend(row)
    for c in DD.CASE_REFS:
        parts.extend(c)
    parts.extend(DD.EXTENSIONS)
    for name, prof in DD.PEER_PROFILES:
        parts.extend([name, prof])
    for k, v in DD.CHECK_LISTS.items():
        parts.append(k); parts.extend(v)
    parts.extend(DD.DEEP_OUTLOOK)
    for t, paras in DD.FOCUS_TOPICS:
        parts.append(t); parts.extend(paras)
    for a, b, lines in DD.REG_QUOTES:
        parts.extend([a, b]); parts.extend(lines)
    parts.extend(DD.ROADMAP["cols"])
    for row in DD.ROADMAP["rows"]:
        parts.extend(row)
    parts.extend(DD.GAP_MATRIX["cols"])
    for row in DD.GAP_MATRIX["rows"]:
        parts.extend(row)
    for t, body in DD.CASE_INTERP:
        parts.extend([t, body])
    for row in DD.TIMELINE:
        parts.extend(row)
    parts.extend(DD.GOV)
    parts.extend(["六大领域深度分析", "行业水位调研", "违规与合规案例参考", "关联延伸",
                  "深度情景展望", "深度分析版", "合规差距自评估矩阵", "案例深度解读",
                  "监管关键节点与合规日历", "合规治理与组织职责建议", "90 天合规落地路线图",
                  "分领域合规自查清单", "同业深度画像", "监管法规条文摘录", "合规焦点专题",
                  "里程碑", "交付物", "阶段", "时间窗", "重点任务", "合规维度",
                  "行业合规水位对比矩阵", "食安快检准入", "价格展示合规", "领先", "达标", "待提升",
                  "图6：朴朴超市合规风险热力矩阵（按风险主题 × 业务环节）"])
    return "".join(parts)


src_simple = collect_chars(include_deep=False)
src_deep = collect_chars(include_deep=True)
src_cjk = sorted(set(c for c in src_deep if '\u4e00' <= c <= '\u9fff'))
cjk_simple = sorted(set(c for c in src_simple if '\u4e00' <= c <= '\u9fff'))

# ---------- 2. 收集数据里的所有链接 ----------
data_urls = []
for _, items in MOD.DATA["policy"]:
    for it in items:
        data_urls.append(it.get("url", ""))
data_urls.extend(r[5] for r in MOD.DATA["penalties"])
data_urls.extend(c[5] for c in DD.CASE_REFS)

def is_root(u):
    return u.rstrip("/").count("/") <= 2 and "beian.cac.gov.cn" not in u

root_urls = [u for u in data_urls if is_root(u)]

# ---------- 3. 逐份 PDF 检测 ----------
targets = [
    ("简版", MOD.META["filename"], 0, cjk_simple, False),
    ("深度版", MOD.META["filename"].replace(".pdf", "_深度分析版.pdf"), 30, src_cjk, True),
]

all_ok = True
print("=" * 68)
print(f"质量门禁检测 · {MOD.META['title']} {MOD.META['date_str']}")
print("=" * 68)
print(f"报告用中文字符总数 : {len(src_cjk)}")
print(f"数据链接总数       : {len(data_urls)}")
print(f"官网首页根链接     : {len(root_urls)}  {'✅ 通过' if not root_urls else '❌ 未通过'}")
for u in root_urls:
    print(f"    违规: {u}")
print("-" * 68)

for label, fname, min_pages, need_chars, is_deep in targets:
    path = os.path.join(OUT_DIR, fname)
    if not os.path.exists(path):
        print(f"[{label}] ❌ 文件不存在: {path}")
        all_ok = False
        continue
    reader = PdfReader(path)
    n_pages = len(reader.pages)
    text = "".join((p.extract_text() or "") for p in reader.pages)
    pdf_cjk = set(c for c in text if '\u4e00' <= c <= '\u9fff')

    # 缺字：仅比对本版实际用到的字符集
    missing = [c for c in need_chars if c not in pdf_cjk]
    tofu = text.count('\ufffd')

    # 可点击链接
    n_links, uris = 0, []
    for p in reader.pages:
        if "/Annots" in p:
            for a in p["/Annots"]:
                obj = a.get_object()
                if obj.get("/Subtype") == "/Link" and "/A" in obj:
                    act = obj["/A"].get_object()
                    if act.get("/S") == "/URI":
                        n_links += 1
                        uris.append(str(act.get("/URI", "")))
    pdf_roots = [u for u in uris if is_root(u)]

    # 章节齐全
    secs = {s: (s in text) for s in MOD.META["sections"].values()}

    # 图示检测：本引擎图示为矢量 Table 绘制（非位图），故按图注 + 内容特征词核验渲染
    n_figs, fig_detail = 0, []
    fig_markers = {
        "行业合规水位对比矩阵": "行业合规水位对比矩阵",
        "食安快检准入 SOP": "供应商到货",
        "价格展示合规样例": "价格展示合规",
        "六大合规领域要点框架": "六大合规领域要点框架",
        "监管要求对照自检矩阵": "监管要求对照自检矩阵",
    }
    if is_deep:
        for name, marker in fig_markers.items():
            hit = marker in text
            n_figs += 1 if hit else 0
            fig_detail.append(f"{name}{'✅' if hit else '❌'}")
    else:
        n_imgs = 0
        for p in reader.pages:
            try:
                res = p.get("/Resources", {})
                if "/XObject" in res:
                    for xo in res["/XObject"].get_object().values():
                        if xo.get_object().get("/Subtype") == "/Image":
                            n_imgs += 1
            except Exception:
                pass
        n_figs, fig_detail = n_imgs, [f"位图 {n_imgs} 张"]

    ok = (not missing) and (tofu == 0) and (not pdf_roots) and all(secs.values()) \
         and n_pages >= min_pages
    if is_deep:
        # 阈值随 fig_markers 自动同步，避免新增图示后忘记改数字导致误判
        ok = ok and n_figs >= len(fig_markers)
    all_ok = all_ok and ok

    print(f"[{label}] {fname}")
    print(f"  页数            : {n_pages}  (要求 ≥{min_pages}) {'✅' if n_pages >= min_pages else '❌'}")
    print(f"  PDF 中文字符数  : {len(pdf_cjk)}")
    print(f"  缺字检测 (P0)   : {len(missing)} {'✅ 通过' if not missing else '❌ 未通过'}")
    if missing:
        print(f"      缺失: {''.join(missing[:80])}")
    print(f"  tofu 替换字符   : {tofu} {'✅' if tofu == 0 else '❌'}")
    print(f"  可点击链接      : {n_links} 条")
    print(f"  PDF 内根链接    : {len(pdf_roots)} {'✅' if not pdf_roots else '❌ ' + str(pdf_roots)}")
    print(f"  章节齐全        : {'✅' if all(secs.values()) else '❌ ' + str([k for k, v in secs.items() if not v])}")
    if is_deep:
        print(f"  图示渲染        : {n_figs}/{len(fig_markers)}  {' ✅ '.join(fig_detail)}")
    else:
        print(f"  热力矩阵位图    : {fig_detail[0]}")
    print(f"  >>> {'通过' if ok else '未通过'}")
    print("-" * 68)

# ---------- 4. DOCX 检测 ----------
try:
    from docx import Document
    for label, fname, need_chars in [("简版", MOD.META["filename"].replace(".pdf", ".docx"), cjk_simple),
                                     ("深度版", MOD.META["filename"].replace(".pdf", "_深度分析版.docx"), src_cjk)]:
        p = os.path.join(OUT_DIR, fname)
        if not os.path.exists(p):
            print(f"[DOCX {label}] ❌ 文件不存在")
            all_ok = False
            continue
        doc = Document(p)
        t = "\n".join(par.text for par in doc.paragraphs)
        for tb in doc.tables:
            for row in tb.rows:
                for cell in row.cells:
                    t += "\n" + cell.text
        miss = [c for c in need_chars if c not in set(t)]
        tofu = t.count('\ufffd')
        print(f"[DOCX {label}] 字符 {len(t)} · 缺字 {len(miss)} {'✅' if not miss else '❌ ' + ''.join(miss[:60])} · tofu {tofu} {'✅' if tofu == 0 else '❌'}")
        all_ok = all_ok and (not miss) and (tofu == 0)
except ImportError:
    print("python-docx 未安装，跳过 DOCX 检测")

print("=" * 68)
print("总体结论:", "✅ 全部通过，允许交付" if all_ok else "❌ 存在未通过项，需修复")
sys.exit(0 if all_ok else 1)
